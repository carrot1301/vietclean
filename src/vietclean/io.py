"""Lightweight, offline input/output helpers for tabular data.

Only :mod:`pandas` is required.  Excel, Parquet and non-SQLite database
support is enabled when their usual optional dependencies are installed.
"""

from __future__ import annotations

import codecs
import csv
import json
import os
import re
import sqlite3
import tempfile
import time
import unicodedata
import warnings
from collections.abc import Callable, Iterable, Sequence
from contextlib import contextmanager, nullcontext
from dataclasses import dataclass, field
from decimal import Decimal
from pathlib import Path
from typing import Any, IO, Iterator, Literal, Mapping, Optional, Tuple, Union, cast
from urllib.parse import unquote

import pandas as pd

from .exceptions import ResourceLimitError, UnsafeInputError, UnsupportedFormatError
from .security.formulas import FormulaMode, protect_spreadsheet_formulas
from .security.resources import ResourceLimits, check_input_file, enforce_frame_limits
from .security.sql import validate_read_query


PathLike = Union[str, os.PathLike]
IfExists = Literal["fail", "replace", "append"]
HeaderRow = Union[int, None, Literal["infer"]]
SheetSelector = Union[str, int, Sequence[Union[str, int]], Literal["all"]]
CancelCallback = Callable[[], bool]
CheckpointCallback = Callable[["ChunkCheckpoint"], None]


@dataclass(frozen=True, slots=True)
class CSVFormat:
    """Detected delimited-text format; callers can override every field."""

    encoding: str
    delimiter: str
    quotechar: str = '"'
    doublequote: bool = True
    escapechar: str | None = None
    quoting: int = csv.QUOTE_MINIMAL


@dataclass(frozen=True, slots=True)
class ChunkCheckpoint:
    """Progress emitted at deterministic chunk boundaries without row contents."""

    phase: Literal["read", "write"]
    chunk_index: int
    rows_processed: int


@dataclass(frozen=True, slots=True)
class ExcelTableRegion:
    """One rectangular non-empty region using one-based Excel coordinates."""

    sheet_name: str
    table_index: int
    start_row: int
    end_row: int
    start_column: int
    end_column: int


@dataclass(frozen=True, slots=True)
class ExcelTable:
    """A detected Excel table together with its bounded DataFrame."""

    region: ExcelTableRegion
    frame: pd.DataFrame


@dataclass(frozen=True, slots=True)
class ExcelSheetInfo:
    """Non-sensitive workbook metadata used for selection and safety warnings."""

    name: str
    index: int
    max_row: int
    max_column: int
    merged_range_count: int
    formula_cell_count: int


@dataclass(frozen=True, slots=True)
class ExcelWorkbookInfo:
    """Bounded metadata for one workbook; never evaluates formulas or VBA."""

    sheets: tuple[ExcelSheetInfo, ...]
    has_macros: bool


@dataclass(frozen=True, slots=True)
class RejectedRow:
    """One local CSV parser rejection; raw fields are never logged automatically."""

    reason: str
    fields: tuple[str, ...]
    line_number: int | None = None
    end_line_number: int | None = None
    record_number: int | None = None


@dataclass(slots=True)
class ReadDiagnostics:
    """Aggregate loader diagnostics plus opt-in locally exportable rejects."""

    encoding: str | None = None
    delimiter: str | None = None
    rejected_rows: list[RejectedRow] = field(default_factory=list)
    rejected_bytes: int = 0
    source_records: int = 0
    accepted_records: int = 0

    def rejects_frame(self) -> pd.DataFrame:
        records = [
            {
                "reject_number": index,
                "source_line": rejected.line_number,
                "source_line_end": rejected.end_line_number,
                "source_record": rejected.record_number,
                "reason": rejected.reason,
                "fields_json": json.dumps(rejected.fields, ensure_ascii=False),
            }
            for index, rejected in enumerate(self.rejected_rows, start=1)
        ]
        return pd.DataFrame.from_records(
            records,
            columns=[
                "reject_number",
                "source_line",
                "source_line_end",
                "source_record",
                "reason",
                "fields_json",
            ],
        )

_DELIMITED_SUFFIXES = {".csv", ".tsv", ".tab", ".txt"}
_EXCEL_SUFFIXES = {".xlsx", ".xls", ".xlsm", ".xlsb"}
_PARQUET_SUFFIXES = {".parquet", ".pq"}
_SQLITE_SUFFIXES = {".db", ".db3", ".sqlite", ".sqlite3"}
_URI_RE = re.compile(r"^[A-Za-z][A-Za-z0-9+.-]*://")
_WINDOWS_DRIVE_RE = re.compile(r"^[A-Za-z]:[\\/]")
_CANDIDATE_DELIMITERS = ",;\t|"
_CP1258_TONE_MARKS = {"\u0300", "\u0301", "\u0303", "\u0309", "\u0323"}
_VIETNAMESE_BASES = set(
    "aAeEiIoOuUyY"
    "\u0103\u0102\u00e2\u00c2\u00ea\u00ca\u00f4\u00d4"
    "\u01a1\u01a0\u01b0\u01af"
)
_VIETNAMESE_LETTERS = set(
    "\u0103\u0102\u00e2\u00c2\u0111\u0110\u00ea\u00ca\u00f4\u00d4"
    "\u01a1\u01a0\u01b0\u01af"
    "\u00e0\u00c0\u00e1\u00c1\u00e3\u00c3\u1ea3\u1ea2\u1ea1\u1ea0"
    "\u1eb1\u1eb0\u1eaf\u1eae\u1eb5\u1eb4\u1eb3\u1eb2\u1eb7\u1eb6"
    "\u1ea7\u1ea6\u1ea5\u1ea4\u1eab\u1eaa\u1ea9\u1ea8\u1ead\u1eac"
    "\u00e8\u00c8\u00e9\u00c9\u1ebd\u1ebc\u1ebb\u1eba\u1eb9\u1eb8"
    "\u1ec1\u1ec0\u1ebf\u1ebe\u1ec5\u1ec4\u1ec3\u1ec2\u1ec7\u1ec6"
    "\u00ec\u00cc\u00ed\u00cd\u0129\u0128\u1ec9\u1ec8\u1ecb\u1eca"
    "\u00f2\u00d2\u00f3\u00d3\u00f5\u00d5\u1ecf\u1ece\u1ecd\u1ecc"
    "\u1ed3\u1ed2\u1ed1\u1ed0\u1ed7\u1ed6\u1ed5\u1ed4\u1ed9\u1ed8"
    "\u1edd\u1edc\u1edb\u1eda\u1ee1\u1ee0\u1edf\u1ede\u1ee3\u1ee2"
    "\u00f9\u00d9\u00fa\u00da\u0169\u0168\u1ee7\u1ee6\u1ee5\u1ee4"
    "\u1eeb\u1eea\u1ee9\u1ee8\u1eef\u1eee\u1eed\u1eec\u1ef1\u1ef0"
    "\u1ef3\u1ef2\u00fd\u00dd\u1ef9\u1ef8\u1ef7\u1ef6\u1ef5\u1ef4"
)


class _OperationGuard:
    """Cooperative timeout/cancellation guard for long-running boundaries."""

    def __init__(
        self,
        timeout_seconds: float | None,
        cancel_callback: CancelCallback | None,
    ) -> None:
        if timeout_seconds is not None:
            if isinstance(timeout_seconds, bool) or not isinstance(
                timeout_seconds, (int, float)
            ):
                raise TypeError("timeout_seconds phải là số hoặc None.")
            if timeout_seconds <= 0:
                raise ValueError("timeout_seconds phải lớn hơn 0.")
        if cancel_callback is not None and not callable(cancel_callback):
            raise TypeError("cancel_callback phải callable hoặc None.")
        self.deadline = (
            time.monotonic() + float(timeout_seconds)
            if timeout_seconds is not None
            else None
        )
        self.cancel_callback = cancel_callback

    def check(self) -> None:
        if self.cancel_callback is not None and self.cancel_callback():
            raise InterruptedError("[VC-CANCEL-001] Thao tác đã được yêu cầu hủy.")
        if self.deadline is not None and time.monotonic() >= self.deadline:
            raise TimeoutError("[VC-TIMEOUT-001] Thao tác vượt thời gian cho phép.")


def detect_encoding(path: PathLike, sample_size: int = 128 * 1024) -> str:
    """Detect a practical text encoding without third-party detectors.

    UTF byte-order marks and valid UTF-8 are preferred.  For legacy files,
    CP1258 is selected only when the decoded sample has convincing Vietnamese
    evidence; otherwise CP1252 and finally Latin-1 are used.
    """

    file_path = Path(path)
    with file_path.open("rb") as stream:
        sample = stream.read(sample_size)
        reached_eof = not stream.read(1)

    if sample.startswith(b"\xef\xbb\xbf"):
        return "utf-8-sig"
    if sample.startswith((b"\xff\xfe", b"\xfe\xff")):
        # Common for Excel's "Unicode Text" exports.
        return "utf-16"
    if not sample:
        return "utf-8"

    try:
        # An incremental decoder avoids a false negative when the bounded
        # sample ends halfway through an otherwise valid multi-byte character.
        codecs.getincrementaldecoder("utf-8")(errors="strict").decode(
            sample, final=reached_eof
        )
        return "utf-8"
    except UnicodeDecodeError:
        pass

    cp1258_text = _try_decode(sample, "cp1258")
    cp1252_text = _try_decode(sample, "cp1252")
    if cp1258_text is not None and _looks_like_cp1258_vietnamese(
        cp1258_text, cp1252_text
    ):
        return "cp1258"
    if cp1252_text is not None:
        return "cp1252"
    return "latin1"


def detect_delimiter(
    path: PathLike,
    *,
    encoding: str = "auto",
    sample_size: int = 64 * 1024,
) -> str:
    """Detect one of comma, semicolon, tab or pipe in a delimited text file."""

    file_path = Path(path)
    actual_encoding = (
        detect_encoding(file_path) if encoding.lower() == "auto" else encoding
    )
    with file_path.open(
        "r", encoding=actual_encoding, errors="strict", newline=""
    ) as stream:
        sample = stream.read(sample_size)

    non_empty_lines = [line for line in sample.splitlines() if line.strip()]
    sniff_sample = "\n".join(non_empty_lines[:50])
    if sniff_sample:
        try:
            return csv.Sniffer().sniff(
                sniff_sample, delimiters=_CANDIDATE_DELIMITERS
            ).delimiter
        except csv.Error:
            guessed = _consistent_delimiter(non_empty_lines[:50])
            if guessed is not None:
                return guessed

    return "\t" if file_path.suffix.lower() in {".tsv", ".tab", ".txt"} else ","


def detect_csv_format(
    path: PathLike,
    *,
    encoding: str = "auto",
    delimiter: str | None = None,
    sample_size: int = 64 * 1024,
) -> CSVFormat:
    """Detect encoding, delimiter and quote settings from a bounded sample.

    Detection is deliberately conservative. If :class:`csv.Sniffer` cannot
    prove a dialect, VietClean keeps standard CSV quoting and only falls back
    to its deterministic delimiter heuristic.
    """

    file_path = Path(path)
    actual_encoding = detect_encoding(file_path) if encoding.lower() == "auto" else encoding
    with file_path.open(
        "r", encoding=actual_encoding, errors="strict", newline=""
    ) as stream:
        sample = stream.read(sample_size)
    non_empty_lines = [line for line in sample.splitlines() if line.strip()]
    sniff_sample = "\n".join(non_empty_lines[:50])
    actual_delimiter = delimiter
    detected: Any = None
    if sniff_sample:
        try:
            detected = csv.Sniffer().sniff(
                sniff_sample,
                delimiters=delimiter if delimiter is not None else _CANDIDATE_DELIMITERS,
            )
        except csv.Error:
            detected = None
    if actual_delimiter is None:
        actual_delimiter = (
            detected.delimiter
            if detected is not None
            else detect_delimiter(file_path, encoding=actual_encoding)
        )
    return CSVFormat(
        encoding=actual_encoding,
        delimiter=actual_delimiter,
        quotechar=(detected.quotechar or '"') if detected is not None else '"',
        doublequote=detected.doublequote if detected is not None else True,
        escapechar=detected.escapechar if detected is not None else None,
        quoting=detected.quoting if detected is not None else csv.QUOTE_MINIMAL,
    )


def read_data(
    source: Union[PathLike, sqlite3.Connection],
    *,
    sheet_name: Union[str, int] = 0,
    table: Optional[str] = None,
    query: Optional[str] = None,
    encoding: str = "auto",
    sep: Optional[str] = None,
    delimiter: Optional[str] = None,
    file_format: Optional[str] = None,
    resource_limits: Union[ResourceLimits, Mapping[str, Any], None] = None,
    allow_macros: bool = False,
    allow_remote_database: bool = False,
    allow_remote_query: bool = False,
    infer_dtypes: bool = False,
    bad_line_policy: Literal["error", "capture"] = "error",
    diagnostics: ReadDiagnostics | None = None,
    header_row: HeaderRow = "infer",
    timeout_seconds: float | None = None,
    cancel_callback: CancelCallback | None = None,
    **options: Any,
) -> pd.DataFrame:
    """Read a local tabular source into a :class:`pandas.DataFrame`.

    ``options`` are passed to the relevant pandas reader.  For example,
    ``header``, ``dtype`` and ``usecols`` work for delimited files, while
    ``params`` can be supplied for a parameterised SQL query.  A SQLite source
    may be a local database path, a ``sqlite:///`` URI, or an existing
    :class:`sqlite3.Connection`.  Other database URIs require SQLAlchemy.

    The function is read-only: it never rewrites or normalises the source.
    """

    _validate_table_query(table, query)
    for name, boolean_value in {
        "allow_macros": allow_macros,
        "allow_remote_database": allow_remote_database,
        "allow_remote_query": allow_remote_query,
        "infer_dtypes": infer_dtypes,
    }.items():
        if not isinstance(boolean_value, bool):
            raise TypeError(f"{name} phải là true hoặc false.")
    if bad_line_policy not in {"error", "capture"}:
        raise ValueError("bad_line_policy phải là error hoặc capture.")
    if bad_line_policy == "capture" and diagnostics is None:
        raise ValueError("bad_line_policy='capture' cần truyền ReadDiagnostics.")
    limits = _coerce_resource_limits(resource_limits)
    guard = _OperationGuard(
        timeout_seconds if timeout_seconds is not None else limits.max_operation_seconds,
        cancel_callback,
    )
    guard.check()
    _validate_header_row(header_row)
    if options.get("chunksize") is not None or options.get("iterator", False):
        raise ValueError(
            "read_data always returns one DataFrame; chunksize/iterator are not supported"
        )
    if sep is not None and delimiter is not None and sep != delimiter:
        raise ValueError("sep and delimiter disagree; provide only one of them")

    if isinstance(source, sqlite3.Connection):
        if file_format is not None and _normalise_format(file_format) != "sqlite":
            raise ValueError("A sqlite3.Connection can only use file_format='sqlite'")
        return enforce_frame_limits(
            _read_sqlite(
                source,
                table=table,
                query=query,
                options=options,
                limits=limits,
                guard=guard,
            ),
            limits,
        )

    kind, value = _classify_location(source, file_format=file_format)
    if kind == "sqlalchemy":
        if not allow_remote_database:
            raise UnsafeInputError(
                "[VC-SQL-001] Remote database bị tắt mặc định để giữ offline-first. "
                "Chỉ bật allow_remote_database=True với credential tối thiểu quyền."
            )
        if query is not None and not allow_remote_query:
            raise UnsafeInputError(
                "[VC-SQL-001] Query tùy chỉnh trên remote database bị tắt mặc định. "
                "Dùng table hoặc bật allow_remote_query=True sau khi review query."
            )
        return enforce_frame_limits(
            _read_sqlalchemy(
                str(value),
                table=table,
                query=query,
                options=options,
                limits=limits,
                guard=guard,
            ),
            limits,
        )
    if kind == "sqlite_uri":
        sqlite_path = _sqlite_path_from_uri(str(value))
        return enforce_frame_limits(
            _read_sqlite_path(
                sqlite_path,
                table=table,
                query=query,
                options=options,
                limits=limits,
                guard=guard,
            ),
            limits,
        )

    path = Path(value)
    if not path.exists():
        raise FileNotFoundError(f"Data source does not exist: {path}")
    if not path.is_file():
        raise ValueError(f"Data source is not a file: {path}")
    check_input_file(path, limits)

    if kind == "delimited":
        if table is not None or query is not None:
            raise ValueError("table/query are only valid for database sources")
        _guard_full_delimited_read(path, limits)
        actual_encoding = (
            detect_encoding(path) if encoding.lower() == "auto" else encoding
        )
        actual_sep = sep if sep is not None else delimiter
        detected_format = detect_csv_format(
            path,
            encoding=actual_encoding,
            delimiter=actual_sep,
        )
        actual_sep = detected_format.delimiter
        if diagnostics is not None:
            diagnostics.encoding = actual_encoding
            diagnostics.delimiter = actual_sep
        if bad_line_policy == "error":
            _scan_delimited_nulls(path, encoding=actual_encoding, guard=guard)
        read_options = dict(options)
        _apply_header_row(read_options, header_row)
        if "dialect" not in read_options:
            read_options.setdefault("quotechar", detected_format.quotechar)
            read_options.setdefault("doublequote", detected_format.doublequote)
            read_options.setdefault("escapechar", detected_format.escapechar)
            read_options.setdefault("quoting", detected_format.quoting)
        # Blank records are data. Keeping them lets the cleaning pipeline audit
        # any configured deletion instead of pandas dropping rows implicitly.
        read_options.setdefault("skip_blank_lines", False)
        if not infer_dtypes and "dtype" not in read_options:
            read_options["dtype"] = "string"
        column_count = _delimited_column_count(
            path,
            encoding=actual_encoding,
            delimiter=actual_sep,
            names=read_options.get("names"),
            parser_options=read_options,
        )
        if column_count > limits.max_columns:
            raise ResourceLimitError(
                f"[VC-LIMIT-001] CSV vượt {limits.max_columns:,} cột."
            )
        if bad_line_policy == "capture":
            if "on_bad_lines" in read_options:
                raise ValueError("Không truyền on_bad_lines khi bad_line_policy='capture'.")
            assert diagnostics is not None
            _validate_capture_options(read_options, actual_sep)
        elif read_options.get("on_bad_lines", "error") != "error":
            raise ValueError(
                "on_bad_lines chỉ được là 'error'; dùng bad_line_policy='capture' "
                "và ReadDiagnostics để không bỏ dòng lỗi âm thầm."
            )
        requested_rows = read_options.get("nrows")
        row_budget = min(
            limits.max_rows,
            limits.max_cells // max(1, column_count),
        )
        safe_probe_rows = row_budget + 1
        if requested_rows is None or int(requested_rows) > safe_probe_rows:
            read_options["nrows"] = safe_probe_rows
        source_context = (
            _filtered_delimited_file(
                path,
                encoding=actual_encoding,
                delimiter=actual_sep,
                diagnostics=diagnostics,
                limits=limits,
                max_data_rows=safe_probe_rows,
                parser_options=read_options,
            )
            if bad_line_policy == "capture"
            else nullcontext(path)
        )
        with source_context as readable_path:
            guard.check()
            frame = pd.read_csv(
                readable_path,
                sep=actual_sep,
                encoding=actual_encoding,
                **read_options,
            )
        guard.check()
        return enforce_frame_limits(frame, limits)

    if kind == "excel":
        if table is not None or query is not None:
            raise ValueError("table/query are only valid for database sources")
        _check_macro_policy(path, allow_macros)
        read_options = dict(options)
        _apply_header_row(read_options, header_row)
        if not infer_dtypes and "dtype" not in read_options:
            read_options["dtype"] = "string"
        requested_rows = read_options.get("nrows")
        if requested_rows is None or int(requested_rows) > limits.max_rows + 1:
            read_options["nrows"] = limits.max_rows + 1
        try:
            guard.check()
            _preflight_excel_selection(path, [sheet_name], limits=limits)
            frame = pd.read_excel(path, sheet_name=sheet_name, **read_options)
        except ImportError as exc:
            raise _excel_dependency_error(path.suffix, "read") from exc
        if not isinstance(frame, pd.DataFrame):
            raise ValueError("sheet_name must select exactly one Excel worksheet")
        guard.check()
        return enforce_frame_limits(frame, limits)

    if kind == "parquet":
        if table is not None or query is not None:
            raise ValueError("table/query are only valid for database sources")
        try:
            _preflight_parquet(path, options=options, limits=limits)
            guard.check()
            frame = pd.read_parquet(path, **options)
            guard.check()
            return enforce_frame_limits(frame, limits)
        except ImportError as exc:
            raise ImportError(
                "Parquet support is optional. Install 'pyarrow' (recommended) "
                "or 'fastparquet'."
            ) from exc

    if kind == "sqlite":
        return enforce_frame_limits(
            _read_sqlite_path(
                path,
                table=table,
                query=query,
                options=options,
                limits=limits,
                guard=guard,
            ),
            limits,
        )

    raise AssertionError(f"Unhandled data source kind: {kind}")


def read_chunks(
    source: PathLike,
    *,
    chunksize: int = 100_000,
    encoding: str = "auto",
    sep: str | None = None,
    delimiter: str | None = None,
    resource_limits: Union[ResourceLimits, Mapping[str, Any], None] = None,
    bad_line_policy: Literal["error", "capture"] = "error",
    diagnostics: ReadDiagnostics | None = None,
    infer_dtypes: bool = False,
    header_row: HeaderRow = "infer",
    start_row: int = 0,
    timeout_seconds: float | None = None,
    cancel_callback: CancelCallback | None = None,
    checkpoint_callback: CheckpointCallback | None = None,
    **options: Any,
) -> Iterator[pd.DataFrame]:
    """Stream a large local CSV/TSV in bounded DataFrame chunks."""

    if isinstance(chunksize, bool) or not isinstance(chunksize, int):
        raise TypeError("chunksize phải là số nguyên.")
    if not isinstance(infer_dtypes, bool):
        raise TypeError("infer_dtypes phải là true hoặc false.")
    if chunksize <= 0:
        raise ValueError("chunksize phải lớn hơn 0.")
    if isinstance(start_row, bool) or not isinstance(start_row, int):
        raise TypeError("start_row phải là số nguyên.")
    if start_row < 0:
        raise ValueError("start_row không được âm.")
    if checkpoint_callback is not None and not callable(checkpoint_callback):
        raise TypeError("checkpoint_callback phải callable hoặc None.")
    if sep is not None and delimiter is not None and sep != delimiter:
        raise ValueError("sep and delimiter disagree; provide only one of them")
    _validate_header_row(header_row)
    if bad_line_policy == "capture" and diagnostics is None:
        raise ValueError("bad_line_policy='capture' cần truyền ReadDiagnostics.")
    path = Path(source)
    kind, value = _classify_location(path, file_format=None)
    if kind != "delimited":
        raise UnsupportedFormatError("read_chunks chỉ hỗ trợ CSV/TSV/TXT cục bộ.")
    limits = _coerce_resource_limits(resource_limits)
    guard = _OperationGuard(
        timeout_seconds if timeout_seconds is not None else limits.max_operation_seconds,
        cancel_callback,
    )
    guard.check()
    check_input_file(Path(value), limits)
    actual_encoding = detect_encoding(path) if encoding.lower() == "auto" else encoding
    detected_format = detect_csv_format(
        path,
        encoding=actual_encoding,
        delimiter=sep if sep is not None else delimiter,
    )
    actual_delimiter = detected_format.delimiter
    if diagnostics is not None:
        diagnostics.encoding = actual_encoding
        diagnostics.delimiter = actual_delimiter
    if bad_line_policy == "error":
        _scan_delimited_nulls(path, encoding=actual_encoding, guard=guard)
    read_options = dict(options)
    _apply_header_row(read_options, header_row)
    if "dialect" not in read_options:
        read_options.setdefault("quotechar", detected_format.quotechar)
        read_options.setdefault("doublequote", detected_format.doublequote)
        read_options.setdefault("escapechar", detected_format.escapechar)
        read_options.setdefault("quoting", detected_format.quoting)
    read_options.setdefault("skip_blank_lines", False)
    if not infer_dtypes and "dtype" not in read_options:
        read_options["dtype"] = "string"
    column_count = _delimited_column_count(
        path,
        encoding=actual_encoding,
        delimiter=actual_delimiter,
        names=read_options.get("names"),
        parser_options=read_options,
    )
    if column_count > limits.max_columns:
        raise ResourceLimitError(f"[VC-LIMIT-001] CSV vượt {limits.max_columns:,} cột.")
    maximum_chunk_rows = max(1, limits.max_cells // max(1, column_count))
    if chunksize > maximum_chunk_rows:
        raise ResourceLimitError(
            "[VC-LIMIT-001] chunksize × số cột vượt max_cells; giảm chunksize."
        )
    if bad_line_policy == "capture":
        if "on_bad_lines" in read_options:
            raise ValueError("Không truyền on_bad_lines khi bad_line_policy='capture'.")
        assert diagnostics is not None
        _validate_capture_options(read_options, actual_delimiter)
    elif bad_line_policy != "error":
        raise ValueError("bad_line_policy phải là error hoặc capture.")
    elif read_options.get("on_bad_lines", "error") != "error":
        raise ValueError(
            "on_bad_lines chỉ được là 'error'; dùng bad_line_policy='capture' "
            "và ReadDiagnostics để không bỏ dòng lỗi âm thầm."
        )

    rows_seen = 0
    source_context = (
        _filtered_delimited_file(
            path,
            encoding=actual_encoding,
            delimiter=actual_delimiter,
            diagnostics=diagnostics,
            limits=limits,
            max_data_rows=min(
                limits.max_rows,
                limits.max_cells // max(1, column_count),
            )
            + 1,
            parser_options=read_options,
        )
        if bad_line_policy == "capture"
        else nullcontext(path)
    )
    with source_context as readable_path:
        reader = pd.read_csv(
            readable_path,
            sep=actual_delimiter,
            encoding=actual_encoding,
            chunksize=chunksize,
            **read_options,
        )
        try:
            emitted_rows = 0
            chunk_index = 0
            rows_to_skip = start_row
            for chunk in reader:
                guard.check()
                rows_seen += len(chunk)
                if rows_seen > limits.max_rows:
                    raise ResourceLimitError(
                        f"[VC-LIMIT-001] Dữ liệu vượt {limits.max_rows:,} dòng."
                    )
                if rows_seen * column_count > limits.max_cells:
                    raise ResourceLimitError(
                        f"[VC-LIMIT-001] Dữ liệu vượt {limits.max_cells:,} ô."
                    )
                if rows_to_skip:
                    skipped = min(rows_to_skip, len(chunk))
                    rows_to_skip -= skipped
                    chunk = chunk.iloc[skipped:]
                if chunk.empty:
                    continue
                chunk_index += 1
                emitted_rows += len(chunk)
                yield enforce_frame_limits(chunk.reset_index(drop=True), limits)
                if checkpoint_callback is not None:
                    checkpoint_callback(
                        ChunkCheckpoint(
                            phase="read",
                            chunk_index=chunk_index,
                            rows_processed=start_row + emitted_rows,
                        )
                    )
                guard.check()
        finally:
            reader.close()


def inspect_excel_workbook(
    source: PathLike,
    *,
    resource_limits: Union[ResourceLimits, Mapping[str, Any], None] = None,
    allow_macros: bool = False,
    timeout_seconds: float | None = None,
    cancel_callback: CancelCallback | None = None,
) -> ExcelWorkbookInfo:
    """Inspect sheet dimensions and formula/merge counts without cell values.

    The workbook is opened read-only and formulas are counted as XML/formula
    records; they are never evaluated. For legacy ``.xls``/``.xlsb`` engines,
    only bounded sheet names are returned because those formats do not expose a
    lightweight common metadata API through pandas.
    """

    if not isinstance(allow_macros, bool):
        raise TypeError("allow_macros phải là true hoặc false.")
    path = Path(source)
    limits = _coerce_resource_limits(resource_limits)
    guard = _OperationGuard(
        timeout_seconds if timeout_seconds is not None else limits.max_operation_seconds,
        cancel_callback,
    )
    check_input_file(path, limits)
    _check_macro_policy(path, allow_macros)
    guard.check()
    suffix = path.suffix.lower()
    if suffix not in _EXCEL_SUFFIXES:
        raise UnsupportedFormatError("inspect_excel_workbook chỉ hỗ trợ file Excel.")
    if suffix in {".xls", ".xlsb"}:
        try:
            with pd.ExcelFile(path) as workbook:
                names = [str(name) for name in workbook.sheet_names]
        except ImportError as exc:
            raise _excel_dependency_error(suffix, "read") from exc
        if len(names) > limits.max_excel_sheets:
            raise ResourceLimitError(
                f"[VC-LIMIT-001] Excel có quá {limits.max_excel_sheets} worksheet."
            )
        return ExcelWorkbookInfo(
            sheets=tuple(
                ExcelSheetInfo(name, index, 0, 0, 0, 0)
                for index, name in enumerate(names)
            ),
            has_macros=_workbook_contains_vba(path),
        )

    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - core dependency in releases
        raise _excel_dependency_error(suffix, "read") from exc
    workbook = load_workbook(
        path,
        read_only=True,
        data_only=False,
        keep_links=False,
    )
    try:
        if len(workbook.sheetnames) > limits.max_excel_sheets:
            raise ResourceLimitError(
                f"[VC-LIMIT-001] Excel có quá {limits.max_excel_sheets} worksheet."
            )
        infos: list[ExcelSheetInfo] = []
        total_declared_cells = 0
        for index, worksheet in enumerate(workbook.worksheets):
            guard.check()
            rows = int(worksheet.max_row or 0)
            columns = int(worksheet.max_column or 0)
            _enforce_excel_dimensions(rows, columns, limits, worksheet.title)
            total_declared_cells += rows * columns
            if total_declared_cells > limits.max_cells:
                raise ResourceLimitError(
                    f"[VC-LIMIT-001] Workbook vượt {limits.max_cells:,} ô khai báo."
                )
            merged_count, formula_count = _count_excel_xml_features(
                path,
                cast(str | None, getattr(worksheet, "_worksheet_path", None)),
            )
            infos.append(
                ExcelSheetInfo(
                    name=worksheet.title,
                    index=index,
                    max_row=rows,
                    max_column=columns,
                    merged_range_count=merged_count,
                    formula_cell_count=formula_count,
                )
            )
        return ExcelWorkbookInfo(
            sheets=tuple(infos),
            has_macros=_workbook_contains_vba(path),
        )
    finally:
        workbook.close()


def read_excel_sheets(
    source: PathLike,
    *,
    sheets: SheetSelector = "all",
    header_row: HeaderRow = "infer",
    resource_limits: Union[ResourceLimits, Mapping[str, Any], None] = None,
    allow_macros: bool = False,
    infer_dtypes: bool = False,
    timeout_seconds: float | None = None,
    cancel_callback: CancelCallback | None = None,
    **options: Any,
) -> dict[str, pd.DataFrame]:
    """Read one, several or all Excel sheets into an ordered mapping.

    Empty internal rows/columns and ``Unnamed`` columns are kept for the
    cleaning plan to decide. Merged data cells are preserved as a value in the
    top-left cell plus missing cells; VietClean never guesses a fill direction.
    """

    _validate_header_row(header_row)
    if not isinstance(allow_macros, bool) or not isinstance(infer_dtypes, bool):
        raise TypeError("allow_macros/infer_dtypes phải là true hoặc false.")
    path = Path(source)
    limits = _coerce_resource_limits(resource_limits)
    guard = _OperationGuard(
        timeout_seconds if timeout_seconds is not None else limits.max_operation_seconds,
        cancel_callback,
    )
    check_input_file(path, limits)
    _check_macro_policy(path, allow_macros)
    info = inspect_excel_workbook(
        path,
        resource_limits=limits,
        allow_macros=allow_macros,
        timeout_seconds=timeout_seconds,
        cancel_callback=cancel_callback,
    )
    selected_names = _select_excel_sheets(
        [sheet.name for sheet in info.sheets],
        sheets,
        limits.max_excel_sheets,
    )
    read_options = dict(options)
    _apply_header_row(read_options, header_row)
    if not infer_dtypes and "dtype" not in read_options:
        read_options["dtype"] = "string"
    result: dict[str, pd.DataFrame] = {}
    total_rows = 0
    total_cells = 0
    for name in selected_names:
        guard.check()
        sheet_info = next(sheet for sheet in info.sheets if sheet.name == name)
        row_budget = min(
            limits.max_rows,
            limits.max_cells // max(1, sheet_info.max_column),
        )
        selected_options = dict(read_options)
        requested_rows = selected_options.get("nrows")
        if requested_rows is None or int(requested_rows) > row_budget + 1:
            selected_options["nrows"] = row_budget + 1
        try:
            frame = pd.read_excel(path, sheet_name=name, **selected_options)
        except ImportError as exc:
            raise _excel_dependency_error(path.suffix, "read") from exc
        enforce_frame_limits(frame, limits)
        total_rows += len(frame)
        total_cells += frame.shape[0] * frame.shape[1]
        if total_rows > limits.max_rows or total_cells > limits.max_cells:
            raise ResourceLimitError(
                "[VC-LIMIT-001] Tổng dữ liệu các sheet vượt giới hạn operation."
            )
        result[name] = frame
    guard.check()
    return result


def detect_excel_tables(
    source: PathLike,
    *,
    sheet_name: str | int = 0,
    resource_limits: Union[ResourceLimits, Mapping[str, Any], None] = None,
    allow_macros: bool = False,
    timeout_seconds: float | None = None,
    cancel_callback: CancelCallback | None = None,
) -> tuple[ExcelTableRegion, ...]:
    """Detect rectangular regions separated by wholly empty rows or columns."""

    path = Path(source)
    if path.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise UnsupportedFormatError(
            "Tự động phát hiện nhiều bảng hiện chỉ hỗ trợ .xlsx/.xlsm."
        )
    limits = _coerce_resource_limits(resource_limits)
    guard = _OperationGuard(
        timeout_seconds if timeout_seconds is not None else limits.max_operation_seconds,
        cancel_callback,
    )
    check_input_file(path, limits)
    _check_macro_policy(path, allow_macros)
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover
        raise _excel_dependency_error(path.suffix, "read") from exc
    workbook = load_workbook(
        path,
        read_only=True,
        data_only=False,
        keep_links=False,
    )
    try:
        name = _select_excel_sheets(workbook.sheetnames, sheet_name, 1)[0]
        worksheet = workbook[name]
        rows = int(worksheet.max_row or 0)
        columns = int(worksheet.max_column or 0)
        _enforce_excel_dimensions(rows, columns, limits, name)
        occupied_by_row: dict[int, set[int]] = {}
        for row in worksheet.iter_rows(
            min_row=1,
            max_row=rows,
            min_col=1,
            max_col=columns,
            values_only=True,
        ):
            guard.check()
            row_number = len(occupied_by_row) + 1
            occupied = {
                index
                for index, value in enumerate(row, start=1)
                if value is not None and (not isinstance(value, str) or value.strip() != "")
            }
            occupied_by_row[row_number] = occupied
        regions: list[ExcelTableRegion] = []
        for row_start, row_end in _contiguous_ranges(
            [row for row, occupied in occupied_by_row.items() if occupied]
        ):
            used_columns = sorted(
                set().union(
                    *(occupied_by_row[row] for row in range(row_start, row_end + 1))
                )
            )
            for column_start, column_end in _contiguous_ranges(used_columns):
                if not any(
                    any(
                        column_start <= column <= column_end
                        for column in occupied_by_row[row]
                    )
                    for row in range(row_start, row_end + 1)
                ):
                    continue
                regions.append(
                    ExcelTableRegion(
                        sheet_name=name,
                        table_index=len(regions),
                        start_row=row_start,
                        end_row=row_end,
                        start_column=column_start,
                        end_column=column_end,
                    )
                )
        return tuple(regions)
    finally:
        workbook.close()


def read_excel_tables(
    source: PathLike,
    *,
    sheet_name: str | int = 0,
    table_indices: Sequence[int] | None = None,
    header_row: int | None = 0,
    resource_limits: Union[ResourceLimits, Mapping[str, Any], None] = None,
    allow_macros: bool = False,
    infer_dtypes: bool = False,
    timeout_seconds: float | None = None,
    cancel_callback: CancelCallback | None = None,
    **options: Any,
) -> tuple[ExcelTable, ...]:
    """Read selected detected table regions; header_row is relative to each region."""

    if header_row is not None and (
        isinstance(header_row, bool) or not isinstance(header_row, int) or header_row < 0
    ):
        raise ValueError("header_row phải là số nguyên không âm hoặc None.")
    if not isinstance(infer_dtypes, bool):
        raise TypeError("infer_dtypes phải là true hoặc false.")
    limits = _coerce_resource_limits(resource_limits)
    guard = _OperationGuard(
        timeout_seconds if timeout_seconds is not None else limits.max_operation_seconds,
        cancel_callback,
    )
    regions = detect_excel_tables(
        source,
        sheet_name=sheet_name,
        resource_limits=limits,
        allow_macros=allow_macros,
        timeout_seconds=timeout_seconds,
        cancel_callback=cancel_callback,
    )
    if table_indices is None:
        selected = list(regions)
    else:
        if isinstance(table_indices, (str, bytes)):
            raise TypeError("table_indices phải là sequence số nguyên.")
        selected = []
        for index in table_indices:
            if isinstance(index, bool) or not isinstance(index, int):
                raise TypeError("Mỗi table index phải là số nguyên.")
            try:
                selected.append(regions[index])
            except IndexError as exc:
                raise ValueError(f"Không có table index {index}.") from exc
    result: list[ExcelTable] = []
    total_cells = 0
    try:
        from openpyxl.utils import get_column_letter
    except ImportError as exc:  # pragma: no cover
        raise _excel_dependency_error(Path(source).suffix, "read") from exc
    for region in selected:
        guard.check()
        height = region.end_row - region.start_row + 1
        width = region.end_column - region.start_column + 1
        total_cells += height * width
        if total_cells > limits.max_cells:
            raise ResourceLimitError(
                "[VC-LIMIT-001] Tổng table regions vượt max_cells."
            )
        usecols = (
            f"{get_column_letter(region.start_column)}:"
            f"{get_column_letter(region.end_column)}"
        )
        selected_options = dict(options)
        if not infer_dtypes and "dtype" not in selected_options:
            selected_options["dtype"] = "string"
        data_rows = (
            height
            if header_row is None
            else height - header_row - 1
        )
        if data_rows < 0:
            raise ValueError(
                f"header_row={header_row} nằm ngoài table region {region.table_index}."
            )
        selected_options.update(
            {
                "skiprows": region.start_row - 1,
                "nrows": data_rows,
                "usecols": usecols,
                "header": header_row,
            }
        )
        frame = pd.read_excel(
            source,
            sheet_name=region.sheet_name,
            **selected_options,
        )
        result.append(ExcelTable(region=region, frame=enforce_frame_limits(frame, limits)))
    guard.check()
    return tuple(result)


def write_data(
    df: pd.DataFrame,
    destination: Union[PathLike, sqlite3.Connection],
    *,
    sheet_name: str = "cleaned",
    table: Optional[str] = None,
    if_exists: IfExists = "fail",
    index: bool = False,
    encoding: str = "utf-8-sig",
    sep: Optional[str] = None,
    delimiter: Optional[str] = None,
    file_format: Optional[str] = None,
    overwrite: bool = False,
    formula_mode: FormulaMode = "sanitize",
    allow_remote_database: bool = False,
    output_root: PathLike | None = None,
    secure_mode: bool = True,
    resource_limits: Union[ResourceLimits, Mapping[str, Any], None] = None,
    timeout_seconds: float | None = None,
    cancel_callback: CancelCallback | None = None,
    **options: Any,
) -> None:
    """Write a DataFrame to a local file or database.

    Local CSV, Excel and Parquet files are written to a sibling temporary file
    and then replaced, so a failed write normally leaves an existing output
    untouched.  Database output requires ``table``; SQLite uses the standard
    library and other database URIs use optional SQLAlchemy.
    """

    if not isinstance(df, pd.DataFrame):
        raise TypeError("df must be a pandas.DataFrame")
    for name, boolean_value in {
        "index": index,
        "overwrite": overwrite,
        "allow_remote_database": allow_remote_database,
        "secure_mode": secure_mode,
    }.items():
        if not isinstance(boolean_value, bool):
            raise TypeError(f"{name} phải là true hoặc false.")
    if sep is not None and delimiter is not None and sep != delimiter:
        raise ValueError("sep and delimiter disagree; provide only one of them")
    if if_exists not in {"fail", "replace", "append"}:
        raise ValueError("if_exists phải là fail, replace hoặc append.")
    limits = _coerce_resource_limits(resource_limits)
    enforce_frame_limits(df, limits)
    guard = _OperationGuard(
        timeout_seconds if timeout_seconds is not None else limits.max_operation_seconds,
        cancel_callback,
    )
    guard.check()

    if isinstance(destination, sqlite3.Connection):
        if file_format is not None and _normalise_format(file_format) != "sqlite":
            raise ValueError("A sqlite3.Connection can only use file_format='sqlite'")
        selected_table = _require_table(table)
        _write_sqlite(
            df,
            destination,
            table=selected_table,
            if_exists=if_exists,
            index=index,
            options=options,
        )
        return

    kind, value = _classify_location(destination, file_format=file_format)
    if kind == "sqlalchemy":
        if not allow_remote_database:
            raise UnsafeInputError(
                "[VC-SQL-001] Ghi remote database bị tắt mặc định. "
                "Chỉ bật allow_remote_database=True khi đích đã được phê duyệt."
            )
        selected_table = _require_table(table)
        _write_sqlalchemy(
            df,
            str(value),
            table=selected_table,
            if_exists=if_exists,
            index=index,
            options=options,
            guard=guard,
        )
        return
    if kind == "sqlite_uri":
        selected_table = _require_table(table)
        sqlite_path = _sqlite_path_from_uri(str(value))
        _write_sqlite_path(
            df,
            sqlite_path,
            table=selected_table,
            if_exists=if_exists,
            index=index,
            options=options,
            output_root=output_root,
            secure_mode=secure_mode,
            guard=guard,
            limits=limits,
        )
        return

    path = Path(value)
    if kind == "sqlite":
        selected_table = _require_table(table)
        _write_sqlite_path(
            df,
            path,
            table=selected_table,
            if_exists=if_exists,
            index=index,
            options=options,
            output_root=output_root,
            secure_mode=secure_mode,
            guard=guard,
            limits=limits,
        )
        return
    if table is not None:
        raise ValueError("table is only valid for database destinations")

    path = validate_output_path(path, output_root=output_root, secure_mode=secure_mode)
    if path.exists() and not overwrite:
        raise FileExistsError(f"Output already exists: {path}. Pass overwrite=True explicitly.")

    if kind == "delimited":
        actual_sep = sep if sep is not None else delimiter
        if actual_sep is None:
            actual_sep = _default_output_delimiter(path, file_format)
        actual_encoding = "utf-8-sig" if encoding.lower() == "auto" else encoding
        protected = protect_spreadsheet_formulas(df, mode=formula_mode).data
        with _atomic_output(
            path,
            overwrite=overwrite,
            output_root=output_root,
            secure_mode=secure_mode,
        ) as temporary_path:
            guard.check()
            protected.to_csv(
                temporary_path,
                sep=actual_sep,
                encoding=actual_encoding,
                index=index,
                **options,
            )
            guard.check()
        return

    if kind == "excel":
        excel_suffix = _excel_output_suffix(path, file_format)
        if excel_suffix != ".xlsx":
            # Recent pandas versions can silently create an XLSX zip stream
            # behind an .xls name with some installed writer combinations.
            # Refuse that misleading output; legacy XLS remains readable via xlrd.
            raise _excel_dependency_error(excel_suffix, "write")
        try:
            protected = protect_spreadsheet_formulas(df, mode=formula_mode).data
            with _atomic_output(
                path,
                temporary_suffix=excel_suffix,
                overwrite=overwrite,
                output_root=output_root,
                secure_mode=secure_mode,
            ) as temporary_path:
                guard.check()
                protected.to_excel(
                    temporary_path,
                    sheet_name=sheet_name,
                    index=index,
                    **options,
                )
                guard.check()
        except ImportError as exc:
            raise _excel_dependency_error(excel_suffix, "write") from exc
        return

    if kind == "parquet":
        try:
            with _atomic_output(
                path,
                overwrite=overwrite,
                output_root=output_root,
                secure_mode=secure_mode,
            ) as temporary_path:
                guard.check()
                df.to_parquet(temporary_path, index=index, **options)
                guard.check()
        except ImportError as exc:
            raise ImportError(
                "Parquet support is optional. Install 'pyarrow' (recommended) "
                "or 'fastparquet'."
            ) from exc
        return

    raise AssertionError(f"Unhandled data destination kind: {kind}")


def write_excel_sheets(
    frames: Mapping[str, pd.DataFrame],
    destination: PathLike,
    *,
    index: bool = False,
    overwrite: bool = False,
    formula_mode: FormulaMode = "sanitize",
    output_root: PathLike | None = None,
    secure_mode: bool = True,
    resource_limits: Union[ResourceLimits, Mapping[str, Any], None] = None,
    timeout_seconds: float | None = None,
    cancel_callback: CancelCallback | None = None,
    **options: Any,
) -> None:
    """Atomically write several DataFrames to a macro-free ``.xlsx`` file."""

    for name, value in {
        "index": index,
        "overwrite": overwrite,
        "secure_mode": secure_mode,
    }.items():
        if not isinstance(value, bool):
            raise TypeError(f"{name} phải là true hoặc false.")
    if not isinstance(frames, Mapping) or not frames:
        raise ValueError("frames phải là mapping không rỗng từ sheet name đến DataFrame.")
    path = Path(destination)
    if path.suffix.lower() != ".xlsx":
        raise UnsupportedFormatError(
            "Multi-sheet output chỉ ghi .xlsx; VietClean không bảo toàn macro."
        )
    limits = _coerce_resource_limits(resource_limits)
    if len(frames) > limits.max_excel_sheets:
        raise ResourceLimitError(
            f"[VC-LIMIT-001] Output vượt {limits.max_excel_sheets} sheet."
        )
    guard = _OperationGuard(
        timeout_seconds if timeout_seconds is not None else limits.max_operation_seconds,
        cancel_callback,
    )
    selected_path = validate_output_path(
        path,
        output_root=output_root,
        secure_mode=secure_mode,
    )
    if selected_path.exists() and not overwrite:
        raise FileExistsError(
            f"Output already exists: {selected_path}. Pass overwrite=True explicitly."
        )
    protected_frames: dict[str, pd.DataFrame] = {}
    total_rows = 0
    total_cells = 0
    casefolded_names: set[str] = set()
    for sheet_name, frame in frames.items():
        guard.check()
        _validate_excel_sheet_name(sheet_name)
        folded = sheet_name.casefold()
        if folded in casefolded_names:
            raise ValueError("Excel sheet names không được trùng khi bỏ qua hoa/thường.")
        casefolded_names.add(folded)
        if not isinstance(frame, pd.DataFrame):
            raise TypeError(f"Sheet {sheet_name!r} không phải DataFrame.")
        enforce_frame_limits(frame, limits)
        total_rows += len(frame)
        total_cells += frame.shape[0] * frame.shape[1]
        if total_rows > limits.max_rows or total_cells > limits.max_cells:
            raise ResourceLimitError(
                "[VC-LIMIT-001] Tổng output workbook vượt giới hạn operation."
            )
        protected_frames[sheet_name] = protect_spreadsheet_formulas(
            frame,
            mode=formula_mode,
        ).data
    try:
        with _atomic_output(
            selected_path,
            temporary_suffix=".xlsx",
            overwrite=overwrite,
            output_root=output_root,
            secure_mode=secure_mode,
        ) as temporary_path:
            with pd.ExcelWriter(temporary_path, engine="openpyxl") as writer:
                for sheet_name, frame in protected_frames.items():
                    guard.check()
                    frame.to_excel(
                        writer,
                        sheet_name=sheet_name,
                        index=index,
                        **options,
                    )
            guard.check()
    except ImportError as exc:  # pragma: no cover - openpyxl is core dependency
        raise _excel_dependency_error(".xlsx", "write") from exc


def write_chunks(
    chunks: Iterable[pd.DataFrame],
    destination: PathLike,
    *,
    encoding: str = "utf-8-sig",
    sep: str | None = None,
    delimiter: str | None = None,
    index: bool = False,
    overwrite: bool = False,
    formula_mode: FormulaMode = "sanitize",
    output_root: PathLike | None = None,
    secure_mode: bool = True,
    resource_limits: Union[ResourceLimits, Mapping[str, Any], None] = None,
    timeout_seconds: float | None = None,
    cancel_callback: CancelCallback | None = None,
    checkpoint_callback: CheckpointCallback | None = None,
    **options: Any,
) -> None:
    """Atomically stream transformed chunks to CSV/TSV with stable columns."""

    for name, value in {
        "index": index,
        "overwrite": overwrite,
        "secure_mode": secure_mode,
    }.items():
        if not isinstance(value, bool):
            raise TypeError(f"{name} phải là true hoặc false.")
    if checkpoint_callback is not None and not callable(checkpoint_callback):
        raise TypeError("checkpoint_callback phải callable hoặc None.")
    if sep is not None and delimiter is not None and sep != delimiter:
        raise ValueError("sep and delimiter disagree; provide only one of them")
    if set(options) & {"mode", "header"}:
        raise ValueError("write_chunks tự quản lý mode/header; không truyền hai option này.")
    path = Path(destination)
    kind, _value = _classify_location(path, file_format=None)
    if kind != "delimited":
        raise UnsupportedFormatError("write_chunks chỉ hỗ trợ CSV/TSV/TXT cục bộ.")
    limits = _coerce_resource_limits(resource_limits)
    guard = _OperationGuard(
        timeout_seconds if timeout_seconds is not None else limits.max_operation_seconds,
        cancel_callback,
    )
    selected_path = validate_output_path(
        path,
        output_root=output_root,
        secure_mode=secure_mode,
    )
    if selected_path.exists() and not overwrite:
        raise FileExistsError(
            f"Output already exists: {selected_path}. Pass overwrite=True explicitly."
        )
    actual_separator = sep if sep is not None else delimiter
    if actual_separator is None:
        actual_separator = _default_output_delimiter(selected_path, None)
    actual_encoding = "utf-8-sig" if encoding.lower() == "auto" else encoding
    expected_columns: list[Any] | None = None
    rows_written = 0
    cells_written = 0
    chunks_written = 0
    with _atomic_output(
        selected_path,
        overwrite=overwrite,
        output_root=output_root,
        secure_mode=secure_mode,
    ) as temporary_path:
        for chunk in chunks:
            guard.check()
            if not isinstance(chunk, pd.DataFrame):
                raise TypeError("Mỗi chunk phải là pandas.DataFrame.")
            enforce_frame_limits(chunk, limits)
            columns = list(chunk.columns)
            if expected_columns is None:
                expected_columns = columns
            elif columns != expected_columns:
                raise ValueError("Các chunk phải có cùng cột và cùng thứ tự.")
            rows_written += len(chunk)
            cells_written += chunk.shape[0] * chunk.shape[1]
            if rows_written > limits.max_rows or cells_written > limits.max_cells:
                raise ResourceLimitError(
                    "[VC-LIMIT-001] Tổng chunk output vượt giới hạn operation."
                )
            protected = protect_spreadsheet_formulas(chunk, mode=formula_mode).data
            protected.to_csv(
                temporary_path,
                mode="w" if chunks_written == 0 else "a",
                header=chunks_written == 0,
                sep=actual_separator,
                encoding=actual_encoding,
                index=index,
                **options,
            )
            chunks_written += 1
            if checkpoint_callback is not None:
                checkpoint_callback(
                    ChunkCheckpoint(
                        phase="write",
                        chunk_index=chunks_written,
                        rows_processed=rows_written,
                    )
                )
        if chunks_written == 0:
            raise ValueError("Không có chunk để ghi; output không được tạo.")
        guard.check()


def validate_output_path(
    destination: PathLike,
    *,
    output_root: PathLike | None = None,
    secure_mode: bool = True,
) -> Path:
    """Validate an output against traversal and symlink/reparse-point attacks."""

    if not isinstance(secure_mode, bool):
        raise TypeError("secure_mode phải là true hoặc false.")
    raw_value = os.fspath(destination)
    if isinstance(raw_value, bytes):
        raw_value = os.fsdecode(raw_value)
    if "\x00" in raw_value:
        raise UnsafeInputError("[VC-SAFE-001] Output path chứa null byte.")
    candidate = Path(raw_value)
    lexical = candidate if candidate.is_absolute() else Path.cwd() / candidate
    if _path_is_link_or_reparse(lexical):
        raise UnsafeInputError("[VC-SAFE-001] Không ghi qua symbolic link/reparse point.")
    root_resolved: Path | None = None
    if output_root is not None:
        root_raw = Path(output_root)
        root_lexical = root_raw if root_raw.is_absolute() else Path.cwd() / root_raw
        if _path_is_link_or_reparse(root_lexical):
            raise UnsafeInputError("[VC-SAFE-001] output_root không được là symlink.")
        root_resolved = root_lexical.resolve(strict=False)
        resolved = lexical.resolve(strict=False)
        try:
            resolved.relative_to(root_resolved)
        except ValueError as exc:
            raise UnsafeInputError(
                "[VC-SAFE-001] Output path nằm ngoài output_root được phép."
            ) from exc
        _assert_no_link_components(lexical.parent, stop=root_lexical)
    elif secure_mode:
        _assert_no_link_components(lexical.parent)
    resolved_candidate = lexical.resolve(strict=False)
    parent = resolved_candidate.parent
    if root_resolved is not None and not root_resolved.exists():
        root_resolved.mkdir(parents=True, exist_ok=True)
    parent.mkdir(parents=True, exist_ok=True)
    _assert_no_link_components(parent, stop=root_resolved)
    if resolved_candidate.exists() and not resolved_candidate.is_file():
        raise UnsafeInputError("[VC-SAFE-001] Output path không phải file.")
    return resolved_candidate


def _try_decode(sample: bytes, encoding: str) -> Optional[str]:
    try:
        return sample.decode(encoding, errors="strict")
    except UnicodeDecodeError:
        return None


def _coerce_resource_limits(
    value: Union[ResourceLimits, Mapping[str, Any], None]
) -> ResourceLimits:
    if value is None:
        return ResourceLimits()
    if isinstance(value, ResourceLimits):
        return value
    if isinstance(value, Mapping):
        return ResourceLimits.from_mapping(value)
    raise TypeError("resource_limits phải là ResourceLimits, mapping hoặc None.")


def _validate_header_row(value: HeaderRow) -> None:
    if value == "infer" or value is None:
        return
    if isinstance(value, bool) or not isinstance(value, int):
        raise TypeError("header_row phải là số nguyên, None hoặc 'infer'.")
    if value < 0:
        raise ValueError("header_row không được âm.")


def _apply_header_row(options: dict[str, Any], header_row: HeaderRow) -> None:
    if header_row == "infer":
        return
    if "header" in options and options["header"] != header_row:
        raise ValueError("header và header_row không được mâu thuẫn.")
    options["header"] = header_row


def _select_excel_sheets(
    available: Sequence[str],
    selector: SheetSelector,
    maximum: int,
) -> list[str]:
    if selector == "all":
        selected_values: list[str | int] = list(available)
    elif isinstance(selector, (str, int)) and not isinstance(selector, bool):
        selected_values = [selector]
    elif isinstance(selector, Sequence) and not isinstance(selector, (str, bytes)):
        selected_values = list(selector)
    else:
        raise TypeError("sheets phải là 'all', tên/index hoặc sequence tên/index.")
    if len(selected_values) > maximum:
        raise ResourceLimitError(
            f"[VC-LIMIT-001] Yêu cầu đọc quá {maximum} worksheet."
        )
    selected: list[str] = []
    for value in selected_values:
        if isinstance(value, bool):
            raise TypeError("Sheet index không được là boolean.")
        if isinstance(value, int):
            index = value if value >= 0 else len(available) + value
            if index < 0 or index >= len(available):
                raise ValueError(f"Không có sheet index {value}.")
            name = available[index]
        elif isinstance(value, str):
            if value not in available:
                raise ValueError(f"Không có sheet tên {value!r}.")
            name = value
        else:
            raise TypeError("Mỗi sheet selector phải là tên hoặc index.")
        if name not in selected:
            selected.append(name)
    if not selected:
        raise ValueError("Phải chọn ít nhất một sheet.")
    return selected


def _enforce_excel_dimensions(
    rows: int,
    columns: int,
    limits: ResourceLimits,
    sheet_name: str,
) -> None:
    if rows > limits.max_rows + 1:
        raise ResourceLimitError(
            f"[VC-LIMIT-001] Sheet {sheet_name!r} vượt {limits.max_rows:,} dòng."
        )
    if columns > limits.max_columns:
        raise ResourceLimitError(
            f"[VC-LIMIT-001] Sheet {sheet_name!r} vượt {limits.max_columns:,} cột."
        )
    if rows * columns > limits.max_cells:
        raise ResourceLimitError(
            f"[VC-LIMIT-001] Sheet {sheet_name!r} vượt {limits.max_cells:,} ô."
        )


def _preflight_excel_selection(
    path: Path,
    selectors: Sequence[str | int],
    *,
    limits: ResourceLimits,
) -> None:
    if path.suffix.lower() in {".xls", ".xlsb"}:
        try:
            with pd.ExcelFile(path) as workbook:
                legacy_names = [str(name) for name in workbook.sheet_names]
        except ImportError as exc:
            raise _excel_dependency_error(path.suffix, "read") from exc
        if len(legacy_names) > limits.max_excel_sheets:
            raise ResourceLimitError(
                f"[VC-LIMIT-001] Excel có quá {limits.max_excel_sheets} worksheet."
            )
        _select_excel_sheets(
            legacy_names,
            list(selectors),
            limits.max_excel_sheets,
        )
        warnings.warn(
            "Không thể preflight kích thước từng sheet cho Excel legacy/binary; "
            "giới hạn sẽ được kiểm tra ngay sau khi đọc.",
            UserWarning,
            stacklevel=3,
        )
        return
    if path.suffix.lower() not in {".xlsx", ".xlsm"}:
        return
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover
        raise _excel_dependency_error(path.suffix, "read") from exc
    workbook = load_workbook(
        path,
        read_only=True,
        data_only=False,
        keep_links=False,
    )
    try:
        selected_names = _select_excel_sheets(
            workbook.sheetnames,
            list(selectors),
            limits.max_excel_sheets,
        )
        total_cells = 0
        for name in selected_names:
            worksheet = workbook[name]
            rows = int(worksheet.max_row or 0)
            columns = int(worksheet.max_column or 0)
            _enforce_excel_dimensions(rows, columns, limits, name)
            total_cells += rows * columns
            if total_cells > limits.max_cells:
                raise ResourceLimitError(
                    "[VC-LIMIT-001] Các sheet được chọn vượt max_cells."
                )
    finally:
        workbook.close()


def _workbook_contains_vba(path: Path) -> bool:
    if path.suffix.lower() not in {".xlsx", ".xlsm", ".xlsb"}:
        return False
    from zipfile import BadZipFile, ZipFile

    try:
        with ZipFile(path) as archive:
            return any(
                name.casefold().endswith("vbaproject.bin")
                for name in archive.namelist()
            )
    except BadZipFile as exc:
        raise UnsafeInputError(
            f"[VC-SAFE-001] File Excel không phải ZIP hợp lệ: {path.name}"
        ) from exc


def _check_macro_policy(path: Path, allow_macros: bool) -> None:
    macro_capable = path.suffix.lower() == ".xlsm" or _workbook_contains_vba(path)
    if macro_capable and not allow_macros:
        raise UnsafeInputError(
            "File có macro/VBA bị từ chối mặc định. Dùng allow_macros=True để chỉ "
            "đọc cell; VietClean không chạy, thực thi hoặc bảo toàn macro."
        )
    if macro_capable:
        warnings.warn(
            "Đang đọc workbook ở chế độ dữ liệu; macro/VBA không được chạy hoặc bảo toàn.",
            UserWarning,
            stacklevel=3,
        )


def _count_excel_xml_features(path: Path, worksheet_path: str | None) -> tuple[int, int]:
    if not worksheet_path:
        return 0, 0
    from zipfile import BadZipFile, ZipFile

    member = worksheet_path.lstrip("/")
    try:
        import xml.etree.ElementTree as element_tree

        with ZipFile(path) as archive:
            with archive.open(member) as stream:
                merged = 0
                formulas = 0
                for _event, element in element_tree.iterparse(stream, events=("start",)):
                    local_name = element.tag.rsplit("}", 1)[-1]
                    if local_name == "mergeCell":
                        merged += 1
                    elif local_name == "f":
                        formulas += 1
                    element.clear()
                return merged, formulas
    except (BadZipFile, KeyError, element_tree.ParseError) as exc:
        raise UnsafeInputError(
            f"[VC-SAFE-001] Worksheet XML không hợp lệ trong {path.name}."
        ) from exc


def _contiguous_ranges(values: Sequence[int]) -> list[tuple[int, int]]:
    if not values:
        return []
    result: list[tuple[int, int]] = []
    start = previous = values[0]
    for value in values[1:]:
        if value != previous + 1:
            result.append((start, previous))
            start = value
        previous = value
    result.append((start, previous))
    return result


def _guard_full_delimited_read(path: Path, limits: ResourceLimits) -> None:
    """Reject a full CSV load likely to exceed the configured RAM budget.

    String-heavy pandas frames plus non-mutating cleaning/audit copies commonly
    occupy many times the compressed/raw CSV bytes. The factor is deliberately
    conservative; callers handling larger files should use ``read_chunks`` or
    explicitly raise the memory limit after measuring their workload.
    """

    estimated_bytes = path.stat().st_size * 12
    maximum_bytes = limits.max_memory_mb * 1024 * 1024
    if estimated_bytes > maximum_bytes:
        raise ResourceLimitError(
            "[VC-LIMIT-001] CSV ước tính vượt ngân sách RAM khi đọc toàn bộ; "
            "dùng read_chunks() hoặc tăng max_memory_mb có chủ đích."
        )


def _scan_delimited_nulls(
    path: Path,
    *,
    encoding: str,
    guard: _OperationGuard,
) -> None:
    """Reject NUL characters before pandas' C parser can silently truncate."""

    try:
        with path.open("r", encoding=encoding, errors="strict", newline="") as stream:
            for line_number, line in enumerate(stream, start=1):
                if line_number % 10_000 == 0:
                    guard.check()
                if "\x00" in line:
                    raise UnsafeInputError(
                        "[VC-SAFE-001] CSV chứa null byte tại dòng "
                        f"{line_number}; dùng capture để lưu reject có truy vết."
                    )
    except UnicodeDecodeError as exc:
        raise UnsafeInputError(
            f"[VC-SAFE-001] CSV không giải mã được bằng {encoding}."
        ) from exc
    guard.check()


def _preflight_parquet(
    path: Path,
    *,
    options: Mapping[str, Any],
    limits: ResourceLimits,
) -> None:
    """Use optional engine metadata to reject oversized Parquet before load."""

    rows: int | None = None
    columns: int | None = None
    expanded_bytes: int | None = None
    requested_columns = options.get("columns")
    try:
        import pyarrow.parquet as parquet

        metadata = parquet.ParquetFile(path).metadata
        rows = int(metadata.num_rows)
        columns = int(metadata.num_columns)
        expanded_bytes = sum(
            int(metadata.row_group(index).total_byte_size)
            for index in range(metadata.num_row_groups)
        )
    except ImportError:
        try:
            import fastparquet

            metadata = fastparquet.ParquetFile(path)
            rows = int(metadata.count())
            columns = len(metadata.columns)
        except ImportError:
            # The actual reader raises the actionable optional dependency error.
            return
    if requested_columns is not None:
        if isinstance(requested_columns, (str, bytes)):
            raise TypeError("Parquet columns phải là sequence tên cột.")
        columns = len(requested_columns)
    assert rows is not None and columns is not None
    if rows > limits.max_rows:
        raise ResourceLimitError(
            f"[VC-LIMIT-001] Parquet vượt {limits.max_rows:,} dòng."
        )
    if columns > limits.max_columns:
        raise ResourceLimitError(
            f"[VC-LIMIT-001] Parquet vượt {limits.max_columns:,} cột."
        )
    if rows * max(1, columns) > limits.max_cells:
        raise ResourceLimitError(
            f"[VC-LIMIT-001] Parquet vượt {limits.max_cells:,} ô."
        )
    if expanded_bytes is not None and expanded_bytes > limits.max_memory_mb * 1024 * 1024:
        raise ResourceLimitError(
            "[VC-LIMIT-001] Parquet metadata ước tính vượt ngân sách RAM."
        )


def _delimited_column_count(
    path: Path,
    *,
    encoding: str,
    delimiter: str,
    names: Any = None,
    parser_options: Mapping[str, Any] | None = None,
) -> int:
    """Read at most one logical CSV record to bound shape before pandas."""

    if names is not None and not isinstance(names, str):
        try:
            return max(1, len(names))
        except TypeError:
            pass
    if len(delimiter) != 1:
        # Regex/multi-character separators keep the conservative byte/RAM guard
        # and post-load shape check; stdlib csv only accepts one character.
        return 1
    try:
        with path.open("r", encoding=encoding, errors="strict", newline="") as stream:
            reader = csv.reader(
                stream,
                **_csv_reader_options(delimiter, parser_options or {}),
            )
            header = (parser_options or {}).get("header", "infer")
            target = 0 if header == "infer" else header
            if target is None:
                target = 0
            if isinstance(target, bool) or not isinstance(target, int) or target < 0:
                raise ValueError("header phải là số nguyên không âm, None hoặc 'infer'.")
            first: list[str] = []
            for _ in range(target + 1):
                first = next(reader, [])
    except csv.Error as exc:
        raise UnsafeInputError(f"[VC-SAFE-001] Không đọc được header CSV: {exc}") from exc
    return max(1, len(first))


def _validate_capture_options(options: Mapping[str, Any], delimiter: str) -> None:
    if len(delimiter) != 1:
        raise ValueError("bad_line_policy='capture' chỉ hỗ trợ delimiter một ký tự.")
    unsupported = {
        "comment",
        "lineterminator",
        "skipfooter",
        "skiprows",
    }
    selected = sorted(set(options) & unsupported)
    if selected:
        raise ValueError(
            "bad_line_policy='capture' không hỗ trợ option làm mất ánh xạ dòng nguồn: "
            f"{selected}. Hãy tiền xử lý rõ ràng hoặc dùng bad_line_policy='error'."
        )
    _csv_reader_options(delimiter, options)


def _csv_reader_options(
    delimiter: str,
    options: Mapping[str, Any],
) -> dict[str, Any]:
    """Translate pandas CSV dialect options to the strict stdlib parser."""

    dialect = options.get("dialect")
    selected: dict[str, Any] = {"delimiter": delimiter, "strict": True}
    if dialect is not None:
        try:
            source_dialect = csv.get_dialect(dialect) if isinstance(dialect, str) else dialect
            for key in (
                "quotechar",
                "doublequote",
                "escapechar",
                "quoting",
                "skipinitialspace",
            ):
                selected[key] = getattr(source_dialect, key)
        except (csv.Error, AttributeError) as exc:
            raise ValueError(f"CSV dialect không hợp lệ: {dialect!r}") from exc
    for key in (
        "quotechar",
        "doublequote",
        "escapechar",
        "quoting",
        "skipinitialspace",
    ):
        if key in options:
            selected[key] = options[key]
    if selected.get("quotechar") == "" and selected.get("quoting") != csv.QUOTE_NONE:
        raise ValueError("quotechar rỗng chỉ hợp lệ khi quoting=csv.QUOTE_NONE.")
    return selected


class _TrackingLines:
    """Track physical lines consumed by one logical :mod:`csv` record."""

    def __init__(self, stream: IO[str]) -> None:
        self.stream = stream
        self.line_number = 0
        self.consumed: list[str] = []

    def __iter__(self) -> "_TrackingLines":
        return self

    def __next__(self) -> str:
        line = next(self.stream)
        self.line_number += 1
        self.consumed.append(line)
        return line

    def reset_record(self) -> None:
        self.consumed.clear()


@contextmanager
def _filtered_delimited_file(
    path: Path,
    *,
    encoding: str,
    delimiter: str,
    diagnostics: ReadDiagnostics | None,
    limits: ResourceLimits,
    max_data_rows: int,
    parser_options: Mapping[str, Any],
) -> Iterator[Path]:
    """Yield a temporary CSV containing only structurally valid records.

    Pandas can silently reinterpret the first over-wide record as an index and
    does not reject short records. A small stdlib CSV pass gives both cases the
    same deterministic policy and records their physical source line.
    """

    if diagnostics is None:  # Defensive: public callers are checked earlier.
        raise ValueError("capture diagnostics không được để trống.")
    descriptor, temporary_name = tempfile.mkstemp(prefix=".vietclean-read-", suffix=".csv")
    temporary_path = Path(temporary_name)
    try:
        with path.open("r", encoding=encoding, errors="strict", newline="") as source:
            with os.fdopen(
                descriptor, "w", encoding=encoding, errors="strict", newline=""
            ) as destination:
                tracking = _TrackingLines(source)
                reader_options = _csv_reader_options(delimiter, parser_options)
                reader = csv.reader(tracking, **reader_options)
                writer_options = {
                    key: value
                    for key, value in reader_options.items()
                    if key != "strict"
                }
                writer = csv.writer(destination, **writer_options, lineterminator="\n")
                names = parser_options.get("names")
                expected_fields: int | None = (
                    len(names)
                    if names is not None and not isinstance(names, str)
                    else None
                )
                header = parser_options.get("header", "infer")
                header_record = 0 if header == "infer" else header
                record_number = 0
                data_records_written = 0
                while data_records_written < max_data_rows:
                    tracking.reset_record()
                    start_line = tracking.line_number + 1
                    try:
                        fields = next(reader)
                    except StopIteration:
                        break
                    except csv.Error:
                        record_number += 1
                        diagnostics.source_records += 1
                        _record_rejected_row(
                            diagnostics,
                            ["".join(tracking.consumed)],
                            limits,
                            reason="csv_parse_error",
                            line_number=start_line,
                            end_line_number=tracking.line_number,
                            record_number=record_number,
                        )
                        # csv.reader cannot reliably recover its internal state;
                        # a fresh reader resumes at the next unread physical line.
                        reader = csv.reader(tracking, **reader_options)
                        continue

                    record_number += 1
                    diagnostics.source_records += 1
                    end_line = tracking.line_number
                    if any("\x00" in field for field in fields):
                        _record_rejected_row(
                            diagnostics,
                            fields,
                            limits,
                            reason="embedded_null_byte",
                            line_number=start_line,
                            end_line_number=end_line,
                            record_number=record_number,
                        )
                        continue

                    is_preamble = isinstance(header_record, int) and record_number <= header_record
                    is_header = isinstance(header_record, int) and record_number == header_record + 1
                    if is_header and expected_fields is None and fields:
                        expected_fields = len(fields)
                    if header_record is None and expected_fields is None and fields:
                        expected_fields = len(fields)
                    if (
                        expected_fields is None
                        and not is_preamble
                        and not is_header
                        and fields
                    ):
                        expected_fields = len(fields)
                    if (
                        not is_preamble
                        and not is_header
                        and fields
                        and expected_fields is not None
                        and len(fields) != expected_fields
                    ):
                        _record_rejected_row(
                            diagnostics,
                            fields,
                            limits,
                            line_number=start_line,
                            end_line_number=end_line,
                            record_number=record_number,
                        )
                        continue
                    writer.writerow(fields)
                    diagnostics.accepted_records += 1
                    if not is_preamble and not is_header:
                        data_records_written += 1
        yield temporary_path
    finally:
        try:
            temporary_path.unlink()
        except FileNotFoundError:
            pass


def _record_rejected_row(
    diagnostics: ReadDiagnostics,
    fields: list[str],
    limits: ResourceLimits,
    *,
    reason: str = "field_count_mismatch",
    line_number: int | None = None,
    end_line_number: int | None = None,
    record_number: int | None = None,
) -> None:
    """Capture one rejected CSV row within explicit memory bounds."""

    if len(diagnostics.rejected_rows) >= limits.max_rejected_rows:
        raise ResourceLimitError(
            f"[VC-LIMIT-001] Rejected rows vượt {limits.max_rejected_rows:,} dòng."
        )
    encoded_size = sum(len(field.encode("utf-8", errors="replace")) for field in fields)
    encoded_size += max(0, len(fields) - 1)
    maximum_bytes = limits.max_rejected_bytes_mb * 1024 * 1024
    if diagnostics.rejected_bytes + encoded_size > maximum_bytes:
        raise ResourceLimitError(
            f"[VC-LIMIT-001] Rejected rows vượt {limits.max_rejected_bytes_mb} MB."
        )
    diagnostics.rejected_rows.append(
        RejectedRow(
            reason=reason,
            fields=tuple(fields),
            line_number=line_number,
            end_line_number=end_line_number,
            record_number=record_number,
        )
    )
    diagnostics.rejected_bytes += encoded_size


def _looks_like_cp1258_vietnamese(
    cp1258_text: str, cp1252_text: Optional[str]
) -> bool:
    valid_combining = 0
    invalid_combining = 0
    previous = ""
    for character in cp1258_text:
        if character in _CP1258_TONE_MARKS:
            if previous in _VIETNAMESE_BASES:
                valid_combining += 1
            else:
                invalid_combining += 1
        if not unicodedata.combining(character):
            previous = character

    cp1258_normalised = unicodedata.normalize("NFC", cp1258_text)
    cp1258_letters = sum(char in _VIETNAMESE_LETTERS for char in cp1258_normalised)
    cp1252_letters = 0
    if cp1252_text is not None:
        cp1252_normalised = unicodedata.normalize("NFC", cp1252_text)
        cp1252_letters = sum(char in _VIETNAMESE_LETTERS for char in cp1252_normalised)

    if valid_combining and valid_combining >= invalid_combining:
        return True
    # Avoid labelling ordinary Western CP1252 text as CP1258 based on one byte.
    return cp1258_letters >= 2 and cp1258_letters >= cp1252_letters + 2


def _consistent_delimiter(lines: list[str]) -> Optional[str]:
    if not lines:
        return None
    best: Optional[Tuple[Tuple[int, int, int], str]] = None
    for candidate in _CANDIDATE_DELIMITERS:
        counts = [line.count(candidate) for line in lines]
        present = [count for count in counts if count > 0]
        if not present:
            continue
        consistency = max(present) - min(present)
        score = (len(present), -consistency, sum(present))
        if best is None or score > best[0]:
            best = (score, candidate)
    if best is None or best[0][0] < max(1, len(lines) // 2):
        return None
    return best[1]


def _normalise_format(file_format: str) -> str:
    if not isinstance(file_format, str):
        raise TypeError("file_format must be a string")
    normalised = file_format.strip().lower().lstrip(".")
    aliases = {
        "csv": "delimited",
        "tsv": "delimited",
        "tab": "delimited",
        "txt": "delimited",
        "text": "delimited",
        "excel": "excel",
        "xlsx": "excel",
        "xls": "excel",
        "xlsm": "excel",
        "xlsb": "excel",
        "parquet": "parquet",
        "pq": "parquet",
        "sqlite": "sqlite",
        "sqlite3": "sqlite",
        "db": "sqlite",
        "sql": "sqlite",
    }
    try:
        return aliases[normalised]
    except KeyError as exc:
        supported = "csv, tsv, txt, xlsx, xls, xlsm, xlsb, parquet, sqlite"
        raise ValueError(
            f"Unsupported file_format {file_format!r}. Supported: {supported}"
        ) from exc


def _default_output_delimiter(path: Path, file_format: Optional[str]) -> str:
    if file_format is not None:
        requested = file_format.strip().lower().lstrip(".")
        return "\t" if requested in {"tsv", "tab", "txt", "text"} else ","
    return "\t" if path.suffix.lower() in {".tsv", ".tab", ".txt"} else ","


def _excel_output_suffix(path: Path, file_format: Optional[str]) -> str:
    if file_format is not None:
        requested = file_format.strip().lower().lstrip(".")
        if requested in {"xlsx", "xls", "xlsm", "xlsb"}:
            return f".{requested}"
    if path.suffix.lower() in _EXCEL_SUFFIXES:
        return path.suffix.lower()
    return ".xlsx"


def _classify_location(
    location: PathLike, *, file_format: Optional[str]
) -> Tuple[str, Union[str, os.PathLike]]:
    if isinstance(location, str) and _is_database_uri(location):
        scheme = location.split("://", 1)[0].lower()
        if scheme in {"sqlite", "sqlite+pysqlite"}:
            if file_format is not None and _normalise_format(file_format) != "sqlite":
                raise ValueError("A SQLite URI can only use file_format='sqlite'")
            return "sqlite_uri", location
        if file_format is not None:
            raise ValueError("file_format is not used with a database URI")
        return "sqlalchemy", location

    try:
        path = Path(location)
    except TypeError as exc:
        raise TypeError("source/destination must be a path, URI, or SQLite connection") from exc

    if file_format is not None:
        return _normalise_format(file_format), path

    suffix = path.suffix.lower()
    if suffix in _DELIMITED_SUFFIXES:
        return "delimited", path
    if suffix in _EXCEL_SUFFIXES:
        return "excel", path
    if suffix in _PARQUET_SUFFIXES:
        return "parquet", path
    if suffix in _SQLITE_SUFFIXES:
        return "sqlite", path
    raise UnsupportedFormatError(
        f"Cannot infer data format from {path.name!r}. Use a supported extension "
        "or pass file_format explicitly."
    )


def _is_database_uri(value: str) -> bool:
    # Do not mistake a Windows drive (for example C:\data.csv) for a URI.
    return not _WINDOWS_DRIVE_RE.match(value) and bool(_URI_RE.match(value))


def _sqlite_path_from_uri(uri: str) -> Union[Path, str]:
    remainder = uri.split("://", 1)[1]
    if "?" in remainder or "#" in remainder:
        raise ValueError(
            "SQLite URI query parameters/fragments are not supported; pass a local "
            "database path or an existing sqlite3.Connection instead."
        )
    if remainder in {"", "/", "/:memory:", ":memory:"}:
        return ":memory:"

    # sqlite:///relative.db -> relative.db
    # sqlite:////absolute/path.db -> /absolute/path.db
    if remainder.startswith("//"):
        raw_path = remainder[1:]
    elif remainder.startswith("/"):
        raw_path = remainder[1:]
    else:
        raise ValueError(
            "SQLite URI must use sqlite:///path/to/database.sqlite syntax"
        )
    raw_path = unquote(raw_path)
    if raw_path == ":memory:":
        return ":memory:"
    return Path(raw_path)


def _validate_table_query(table: Optional[str], query: Optional[str]) -> None:
    if table is not None and not isinstance(table, str):
        raise TypeError("table must be a string")
    if query is not None and not isinstance(query, str):
        raise TypeError("query must be a string")
    if table is not None and query is not None:
        raise ValueError("Provide either table or query, not both")
    if table is not None and not str(table).strip():
        raise ValueError("table cannot be empty")
    if query is not None and not str(query).strip():
        raise ValueError("query cannot be empty")


def _require_table(table: Optional[str]) -> str:
    if table is not None and not isinstance(table, str):
        raise TypeError("table must be a string")
    if table is None or not table.strip():
        raise ValueError("table is required for database output")
    return table


def _quote_sqlite_identifier(identifier: str) -> str:
    return '"' + identifier.replace('"', '""') + '"'


def _read_sqlite_path(
    path: Union[Path, str],
    *,
    table: Optional[str],
    query: Optional[str],
    options: Mapping[str, Any],
    limits: ResourceLimits,
    guard: _OperationGuard,
) -> pd.DataFrame:
    if path != ":memory:":
        file_path = Path(path)
        if not file_path.exists():
            raise FileNotFoundError(f"SQLite database does not exist: {file_path}")
        if not file_path.is_file():
            raise ValueError(f"SQLite source is not a file: {file_path}")
        check_input_file(file_path, limits)
        # URI mode=ro protects the source even if a caller accidentally passes
        # a mutating SQL statement as ``query``.
        connection = sqlite3.connect(
            f"{file_path.resolve().as_uri()}?mode=ro", uri=True
        )
    else:
        connection = sqlite3.connect(":memory:")
    try:
        return _read_sqlite(
            connection,
            table=table,
            query=query,
            options=options,
            limits=limits,
            guard=guard,
        )
    finally:
        connection.close()


def _read_sqlite(
    connection: sqlite3.Connection,
    *,
    table: Optional[str],
    query: Optional[str],
    options: Mapping[str, Any],
    limits: ResourceLimits,
    guard: _OperationGuard,
) -> pd.DataFrame:
    _validate_table_query(table, query)
    previous_query_only = int(connection.execute("PRAGMA query_only").fetchone()[0])
    if not previous_query_only:
        connection.execute("PRAGMA query_only = ON")
    interrupted: list[BaseException] = []

    def progress_handler() -> int:
        try:
            guard.check()
        except (InterruptedError, TimeoutError) as exc:
            interrupted.append(exc)
            return 1
        return 0

    connection.set_progress_handler(progress_handler, 10_000)
    try:
        guard.check()
        selected_query = validate_read_query(query) if query is not None else None
        if selected_query is None:
            selected_table = table or _single_sqlite_table(connection)
            selected_query = (
                f"SELECT * FROM {_quote_sqlite_identifier(selected_table)}"
            )
        probe_query = "SELECT * FROM (" + selected_query + ") AS __vietclean_source LIMIT 0"
        params = options.get("params")
        probe = (
            connection.execute(probe_query)
            if params is None
            else connection.execute(probe_query, params)
        )
        column_count = len(probe.description or ())
        if column_count > limits.max_columns:
            raise ResourceLimitError(
                f"[VC-LIMIT-001] Query SQLite vượt {limits.max_columns:,} cột."
            )
        row_budget = min(
            limits.max_rows,
            limits.max_cells // max(1, column_count),
        )
        bounded_query = (
            "SELECT * FROM ("
            + selected_query
            + f") AS __vietclean_source LIMIT {row_budget + 1}"
        )
        try:
            frame = pd.read_sql_query(bounded_query, connection, **dict(options))
        except sqlite3.OperationalError:
            if interrupted:
                raise interrupted[0] from None
            raise
        guard.check()
        return frame
    finally:
        connection.set_progress_handler(None, 0)
        if not previous_query_only:
            connection.execute("PRAGMA query_only = OFF")


def _single_sqlite_table(connection: sqlite3.Connection) -> str:
    rows = connection.execute(
        "SELECT name FROM sqlite_master "
        "WHERE type IN ('table', 'view') AND name NOT LIKE 'sqlite_%' "
        "ORDER BY name"
    ).fetchall()
    names = [str(row[0]) for row in rows]
    if len(names) == 1:
        return names[0]
    if not names:
        raise ValueError("SQLite database contains no readable tables; provide query")
    preview = ", ".join(names[:10])
    if len(names) > 10:
        preview += ", ..."
    raise ValueError(
        f"SQLite database has multiple tables ({preview}); provide table or query"
    )


def _write_sqlite_path(
    frame: pd.DataFrame,
    path: Union[Path, str],
    *,
    table: str,
    if_exists: IfExists,
    index: bool,
    options: Mapping[str, Any],
    output_root: PathLike | None,
    secure_mode: bool,
    guard: _OperationGuard,
    limits: ResourceLimits,
) -> None:
    if path == ":memory:":
        connection = sqlite3.connect(":memory:")
        try:
            _write_sqlite(
                frame,
                connection,
                table=table,
                if_exists=if_exists,
                index=index,
                options=options,
            )
        finally:
            connection.close()
        return

    file_path = validate_output_path(
        Path(path),
        output_root=output_root,
        secure_mode=secure_mode,
    )
    if file_path.exists():
        check_input_file(file_path, limits)
    # Work on a sibling SQLite snapshot and replace only after commit. This
    # preserves an existing database when schema conversion or insertion fails.
    with _atomic_output(
        file_path,
        temporary_suffix=file_path.suffix,
        overwrite=True,
        output_root=output_root,
        secure_mode=secure_mode,
    ) as temporary_path:
        guard.check()
        if file_path.exists():
            source = sqlite3.connect(
                f"{file_path.resolve().as_uri()}?mode=ro",
                uri=True,
            )
            target = sqlite3.connect(temporary_path)
            try:
                source.backup(target)
            finally:
                target.close()
                source.close()
        connection = sqlite3.connect(temporary_path)
        try:
            _write_sqlite(
                frame,
                connection,
                table=table,
                if_exists=if_exists,
                index=index,
                options=options,
            )
        finally:
            connection.close()
        guard.check()


def _write_sqlite(
    frame: pd.DataFrame,
    connection: sqlite3.Connection,
    *,
    table: str,
    if_exists: IfExists,
    index: bool,
    options: Mapping[str, Any],
) -> None:
    try:
        with connection:
            _prepare_sqlite_frame(frame).to_sql(
                table,
                connection,
                if_exists=if_exists,
                index=index,
                **dict(options),
            )
    except Exception:
        connection.rollback()
        raise


def _prepare_sqlite_frame(frame: pd.DataFrame) -> pd.DataFrame:
    """Return SQLite-bindable data without losing ``Decimal`` precision.

    Python's built-in sqlite driver cannot bind :class:`~decimal.Decimal` and
    converting financial values to float would be lossy. Columns containing a
    Decimal are therefore copied and stored as canonical TEXT values. Other
    columns retain their original dtype.
    """

    prepared = frame.copy(deep=False)
    copied = False
    for column in frame.columns:
        series = frame[column]
        has_decimal = bool(series.map(lambda value: isinstance(value, Decimal)).any())
        if not has_decimal:
            continue
        if not copied:
            prepared = frame.copy(deep=True)
            copied = True
        prepared[column] = series.map(
            lambda value: str(value) if isinstance(value, Decimal) else value
        )
    return prepared


def _load_sqlalchemy() -> Any:
    try:
        import sqlalchemy
    except ImportError as exc:
        raise ImportError(
            "Database URIs other than SQLite require the optional 'sqlalchemy' "
            "package and the matching database driver."
        ) from exc
    return sqlalchemy


def _read_sqlalchemy(
    uri: str,
    *,
    table: Optional[str],
    query: Optional[str],
    options: Mapping[str, Any],
    limits: ResourceLimits,
    guard: _OperationGuard,
) -> pd.DataFrame:
    _validate_table_query(table, query)
    if table is None and query is None:
        raise ValueError("table or query is required for a database URI")
    sqlalchemy = _load_sqlalchemy()
    try:
        engine = sqlalchemy.create_engine(uri)
    except (
        ImportError,
        ModuleNotFoundError,
        sqlalchemy.exc.NoSuchModuleError,
    ) as exc:
        raise ImportError(
            "The SQLAlchemy database driver required by this URI is not installed."
        ) from exc
    try:
        guard.check()
        with engine.connect().execution_options(stream_results=True) as connection:
            selected_options = dict(options)
            if query is not None:
                safe_query = validate_read_query(query)
                statement = sqlalchemy.text(safe_query)
                iterator = pd.read_sql_query(
                    statement,
                    connection,
                    chunksize=min(10_000, limits.max_rows + 1),
                    **selected_options,
                )
            else:
                iterator = pd.read_sql_table(
                    str(table),
                    connection,
                    chunksize=min(10_000, limits.max_rows + 1),
                    **selected_options,
                )
            return _consume_sql_chunks(iterator, limits=limits, guard=guard)
    except (ResourceLimitError, InterruptedError, TimeoutError, ValueError):
        raise
    except sqlalchemy.exc.SQLAlchemyError:
        # Deliberately omit URI, driver message and credentials from the public
        # exception. Detailed DB diagnostics belong in a separately secured log.
        raise ConnectionError(
            "[VC-SQL-001] Không thể đọc remote database; chi tiết kết nối đã được ẩn."
        ) from None
    finally:
        engine.dispose()


def _consume_sql_chunks(
    chunks: Iterable[pd.DataFrame],
    *,
    limits: ResourceLimits,
    guard: _OperationGuard,
) -> pd.DataFrame:
    frames: list[pd.DataFrame] = []
    rows_seen = 0
    column_count: int | None = None
    for chunk in chunks:
        guard.check()
        if column_count is None:
            column_count = chunk.shape[1]
            if column_count > limits.max_columns:
                raise ResourceLimitError(
                    f"[VC-LIMIT-001] Query vượt {limits.max_columns:,} cột."
                )
        row_budget = min(
            limits.max_rows,
            limits.max_cells // max(1, column_count),
        )
        remaining_probe = row_budget + 1 - rows_seen
        if remaining_probe <= 0:
            break
        selected = chunk.iloc[:remaining_probe]
        frames.append(selected)
        rows_seen += len(selected)
        if rows_seen > row_budget:
            raise ResourceLimitError(
                "[VC-LIMIT-001] Query vượt giới hạn dòng/ô trước khi load toàn bộ."
            )
    guard.check()
    if not frames:
        return pd.DataFrame()
    return pd.concat(frames, ignore_index=True)


def _write_sqlalchemy(
    frame: pd.DataFrame,
    uri: str,
    *,
    table: str,
    if_exists: IfExists,
    index: bool,
    options: Mapping[str, Any],
    guard: _OperationGuard,
) -> None:
    sqlalchemy = _load_sqlalchemy()
    try:
        engine = sqlalchemy.create_engine(uri)
    except (
        ImportError,
        ModuleNotFoundError,
        sqlalchemy.exc.NoSuchModuleError,
    ) as exc:
        raise ImportError(
            "The SQLAlchemy database driver required by this URI is not installed."
        ) from exc
    try:
        guard.check()
        with engine.begin() as connection:
            frame.to_sql(
                table,
                connection,
                if_exists=if_exists,
                index=index,
                **dict(options),
            )
        guard.check()
    except (InterruptedError, TimeoutError, ValueError):
        raise
    except sqlalchemy.exc.SQLAlchemyError:
        raise ConnectionError(
            "[VC-SQL-001] Không thể ghi remote database; chi tiết kết nối đã được ẩn."
        ) from None
    finally:
        engine.dispose()


def _excel_dependency_error(suffix: str, action: str) -> ImportError:
    suffix = suffix.lower()
    if suffix == ".xls":
        if action == "read":
            detail = "Install 'xlrd' to read legacy .xls files."
        else:
            detail = (
                "Legacy .xls writing is not supported by current pandas releases; "
                "write .xlsx instead."
            )
    elif suffix == ".xlsb":
        if action == "read":
            detail = "Install the optional 'pyxlsb' package to read .xlsb files."
        else:
            detail = "Writing .xlsb is not supported; write .xlsx instead."
    elif suffix == ".xlsm" and action == "write":
        detail = (
            "Macro-enabled output is disabled because VietClean does not preserve VBA; "
            "write .xlsx instead."
        )
    else:
        detail = "Install 'openpyxl' for .xlsx/.xlsm support."
    return ImportError(f"Cannot {action} Excel file. {detail}")


@contextmanager
def _atomic_output(
    destination: Path,
    *,
    temporary_suffix: Optional[str] = None,
    overwrite: bool,
    output_root: PathLike | None,
    secure_mode: bool,
) -> Iterator[Path]:
    destination = validate_output_path(
        destination,
        output_root=output_root,
        secure_mode=secure_mode,
    )
    if destination.exists() and not overwrite:
        raise FileExistsError(
            f"Output already exists: {destination}. Pass overwrite=True explicitly."
        )
    parent_identity = _directory_identity(destination.parent)
    descriptor, temporary_name = tempfile.mkstemp(
        prefix=f".{destination.stem}.",
        suffix=temporary_suffix if temporary_suffix is not None else destination.suffix,
        dir=str(destination.parent),
    )
    os.close(descriptor)
    temporary_path = Path(temporary_name)
    try:
        yield temporary_path
        if _directory_identity(destination.parent) != parent_identity:
            raise UnsafeInputError(
                "[VC-SAFE-001] Output directory đã thay đổi trong khi ghi."
            )
        if _path_is_link_or_reparse(destination):
            raise UnsafeInputError(
                "[VC-SAFE-001] Output bị thay bằng symbolic link/reparse point."
            )
        if destination.exists() and not overwrite:
            raise FileExistsError(
                f"Output appeared during write: {destination}."
            )
        os.replace(temporary_path, destination)
    finally:
        try:
            temporary_path.unlink()
        except FileNotFoundError:
            pass


def _validate_excel_sheet_name(name: str) -> None:
    if not isinstance(name, str) or not name:
        raise ValueError("Excel sheet name phải là chuỗi không rỗng.")
    if len(name) > 31:
        raise ValueError("Excel sheet name không được dài quá 31 ký tự.")
    if any(character in name for character in "[]:*?/\\"):
        raise ValueError(f"Excel sheet name chứa ký tự cấm: {name!r}")
    if name.startswith("'") or name.endswith("'"):
        raise ValueError("Excel sheet name không được bắt đầu/kết thúc bằng dấu nháy đơn.")


def _path_is_link_or_reparse(path: Path) -> bool:
    try:
        metadata = path.lstat()
    except FileNotFoundError:
        return False
    if path.is_symlink():
        return True
    attributes = getattr(metadata, "st_file_attributes", 0)
    return bool(attributes & 0x400)


def _assert_no_link_components(path: Path, *, stop: Path | None = None) -> None:
    absolute = path if path.is_absolute() else Path.cwd() / path
    stop_absolute = None
    if stop is not None:
        stop_absolute = stop if stop.is_absolute() else Path.cwd() / stop
    components: list[Path] = []
    current = absolute
    while True:
        components.append(current)
        if stop_absolute is not None and current == stop_absolute:
            break
        if current.parent == current:
            break
        current = current.parent
    for component in reversed(components):
        if component.exists() and _path_is_link_or_reparse(component):
            raise UnsafeInputError(
                f"[VC-SAFE-001] Output path đi qua symlink/reparse point: {component.name}"
            )


def _directory_identity(path: Path) -> tuple[int, int]:
    metadata = path.stat()
    return int(metadata.st_dev), int(metadata.st_ino)


__all__ = [
    "CSVFormat",
    "ChunkCheckpoint",
    "ExcelSheetInfo",
    "ExcelTable",
    "ExcelTableRegion",
    "ExcelWorkbookInfo",
    "ReadDiagnostics",
    "RejectedRow",
    "detect_csv_format",
    "detect_delimiter",
    "detect_encoding",
    "detect_excel_tables",
    "inspect_excel_workbook",
    "read_chunks",
    "read_data",
    "read_excel_sheets",
    "read_excel_tables",
    "validate_output_path",
    "write_chunks",
    "write_data",
    "write_excel_sheets",
]
